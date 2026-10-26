from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/research/repos.md'
OUTPUT = ROOT / 'docs/research/rencana-dan-repo-motionkit.docx'


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    element = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        item = OxmlElement(f'w:{edge}')
        item.set(qn('w:val'), 'single')
        item.set(qn('w:sz'), '4')
        item.set(qn('w:color'), 'D9D9D9')
        element.append(item)
    tc_pr.append(element)


def cell_text(cell, value, bold=False, color='222222'):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(value)
    r.bold = bold
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    borders(cell)
    margins = OxmlElement('w:tcMar')
    for side in ('top', 'start', 'bottom', 'end'):
        part = OxmlElement(f'w:{side}')
        part.set(qn('w:w'), '110')
        part.set(qn('w:type'), 'dxa')
        margins.append(part)
    cell._tc.get_or_add_tcPr().append(margins)


def link(cell, url):
    cell.text = ''
    p = cell.paragraphs[0]
    part = p.part
    rid = part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), rid)
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '245D77')
    props.append(color)
    run.append(props)
    t = OxmlElement('w:t')
    t.text = 'Buka repo'
    run.append(t)
    h.append(run)
    p._p.append(h)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    borders(cell)


def add_table(doc, headings, rows, widths):
    table = doc.add_table(rows=1, cols=len(headings))
    table.autofit = False
    for i, title in enumerate(headings):
        table.columns[i].width = Cm(widths[i])
        table.rows[0].cells[i].width = Cm(widths[i])
        cell_text(table.rows[0].cells[i], title, True, 'FFFFFF')
        shade(table.rows[0].cells[i], '293D4B')
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    for row_i, values in enumerate(rows):
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for i, value in enumerate(values):
            row.cells[i].width = Cm(widths[i])
            if headings[i] == 'Tautan':
                link(row.cells[i], value)
            else:
                cell_text(row.cells[i], value)
            if row_i % 2:
                shade(row.cells[i], 'F2F6F7')
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def item(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(text)


doc = Document()
section = doc.sections[0]
section.page_height = Cm(29.7)
section.page_width = Cm(21)
section.top_margin = Cm(2)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2)
section.right_margin = Cm(2)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(9.5)
styles['Normal'].paragraph_format.space_after = Pt(5)
styles['Normal'].paragraph_format.line_spacing = 1.12
for name, size, before, after in [('Title', 20, 0, 11), ('Heading 1', 13, 12, 6), ('Heading 2', 10.5, 8, 4)]:
    st = styles[name]
    st.font.name = 'Aptos Display' if name == 'Title' else 'Aptos'
    st.font.size = Pt(size)
    st.font.bold = name != 'Title'
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True
    if name == 'Title':
        ppr = st._element.get_or_add_pPr()
        for old in ppr.findall(qn('w:pBdr')):
            ppr.remove(old)
        border = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'nil')
        border.append(bottom)
        ppr.append(border)

title = doc.add_paragraph('Rencana dan Status Riset MotionKit', style='Title')
doc.add_paragraph('20 Oktober 2026  |  Ringkasan kerja dan katalog repo')
doc.add_paragraph('Tujuan kita adalah membuat MotionKit mampu mengubah brief video menjadi adegan yang memakai model, karakter, lingkungan, kamera, cahaya, dan suara yang sesuai. Pekerjaan saat ini berfokus pada bukti adegan pasar kayu Indonesia. Library reusable dan pemilih asset otomatis masih menjadi pekerjaan berikutnya.')

doc.add_heading('Rencana kerja', level=1)
plan = [
    ('1. Fondasi dan katalog', 'Kelompokkan repo menurut fungsi; catat lisensi, format, skala, animasi, dan asal asset.'),
    ('2. Uji asset per jalur', 'Ambil satu contoh kecil untuk model, rig, bangunan, kendaraan, suasana, dan audio. Render still dan catat batasnya.'),
    ('3. Library model dan scene', 'Buat komponen karakter, lapak, barang dagangan, props, kendaraan, serta preset lighting dan SFX yang dapat dipakai ulang.'),
    ('4. Interaksi dan keramaian', 'Uji gerak berjalan, tangan mengambil barang, timbangan, serah terima, kerumunan, dan kamera; cek kontak fisik dan kontinuitas.'),
    ('5. Brief ke video', 'Petakan kebutuhan brief ke komponen, susun script dan storyboard, buat proof shot, periksa still, lalu render film.'),
    ('6. Kualitas dan integrasi', 'Periksa referensi, lisensi, performa, LOD, audio, grid komposisi, dan hasil video sebelum modul dinaikkan ke library inti.'),
]
add_table(doc, ['Tahap', 'Hasil yang dituju'], plan, [4.1, 12.8])
doc.add_paragraph('Delapan jalur bahasan kita: asset dan model 3D; rigging dan animasi; bangunan dan environment; kendaraan; lighting, weather dan post-processing; audio dan SFX; timeline/editor; serta agent workflow, katalog dan QA.')

doc.add_heading('Yang sudah dikerjakan', level=1)
done = [
    ('Katalog dan kontrak', 'Daftar repo, roadmap, matriks integrasi, dan rancangan metadata asset sudah ditulis.', 'Selesai sebagai dokumen kerja'),
    ('GLB dan rig', 'Fox dan CesiumMan dari Khronos diunduh per asset; GLB dirender pada beberapa frame. CesiumMan diberi material netral.', 'Bukti teknis; belum komponen produksi'),
    ('Bangunan', 'GLB Industrial Outpost ModKit CC0 diperiksa dari tiga sudut. Kios kelontong serta lapak pasar kayu dibuat secara prosedural.', 'Bukti bentuk; detail masih perlu peningkatan'),
    ('Pasar dan karakter', 'Lapak ayam, ikan, sayur serta display barang; karakter tanpa wajah berambut dengan lengan IK dua ruas, jari, dan gaya jalan sederhana.', 'Prototipe bergaya, bukan simulasi realistis'),
    ('Interaksi', 'Video 6 detik VendorAction dan VendorWeigh: raih, genggam, angkat, timbang, letakkan; suara sintetis.', 'Video proof tersedia'),
    ('Pasar ramai', 'Video 10 detik FullMarket: enam lapak, enam penjual, dua pembeli, empat pejalan, dua arah gerak.', 'Video proof; keramaian dan stok belum natural'),
    ('Transaksi dekat', 'Komposisi MarketExchange 7 detik dibuat; still pada momen serah diperiksa dan posisi kamera serta kantong direvisi.', 'Sedang disempurnakan; video final belum diverifikasi'),
]
add_table(doc, ['Area', 'Bukti di workspace', 'Status'], done, [3.0, 9.6, 4.3])
doc.add_paragraph('Semua contoh pasar adalah interpretasi stylized dari foto referensi pengguna. Susunan pasar, pakaian, dan gerak detailnya direka untuk proof; belum rekonstruksi lokasi tertentu. Model Khronos dan ModKit memiliki catatan atribusi/lisensi tersendiri.')

doc.add_heading('Repo dan kegunaannya', level=1)
doc.add_paragraph('Status “Kandidat” berarti tercatat di katalog, belum di-clone atau diuji di MotionKit. “Dipakai” berarti sebagian asset atau fungsi sudah diuji; ini tidak berarti seluruh repo diintegrasikan.')
content = SOURCE.read_text(encoding='utf-8')
groups = []
group = None
for line in content.splitlines():
    if line.startswith('## '):
        group = (line[3:], [])
        groups.append(group)
    match = re.match(r'- \[([^]]+)\]\(([^)]+)\) — (.+)', line)
    if match and group:
        group[1].append(match.groups())

used = {'Three.js', 'Khronos glTF Sample Assets', 'ModKit'}
existing = {'pmndrs/postprocessing'}
for title, entries in groups:
    doc.add_heading(title.replace('3D', '3D'), level=2)
    rows = []
    for name, url, purpose in entries:
        status = 'Dipakai sebagian' if name in used else ('Sudah terpasang; belum diuji di jalur ini' if name in existing else 'Kandidat')
        rows.append((name, purpose, status, url))
    add_table(doc, ['Repo', 'Kegunaan untuk MotionKit', 'Status', 'Tautan'], rows, [4.2, 8.4, 2.8, 1.5])

doc.add_heading('Prioritas berikutnya', level=1)
item(doc, 'Selesaikan transaksi close-up: satu kantong berpindah dari penjual ke pembeli, kontak tangan terbaca, lalu render dan periksa MP4.')
item(doc, 'Perbaiki variasi karakter, stok sayur/ayam/ikan, dan kepadatan latar pasar; uji gerak kaki yang tidak meluncur.')
item(doc, 'Pilih satu modul berikutnya secara terukur: lighting dan ambience pasar, atau pipeline asset reusable; buat satu proof yang bisa dibandingkan.')
item(doc, 'Setelah proof memenuhi standar visual, pindahkan komponen yang stabil dari scripts/lab ke library MotionKit dan buat contoh pemanggilan dari brief.')

doc.add_heading('Catatan penggunaan repo', level=1)
doc.add_paragraph('Tidak perlu clone semua repo sekaligus. Sampai tahap ini, asset GLB tertentu diunduh dari Khronos dan ModKit; repo lain terutama dicatat dari riset dan perlu pemeriksaan sumber, lisensi, API, serta hasil render sebelum dipakai. Wadi telah diaudit: implementasi saat ini TypeScript/WDL dan desain rumahnya tidak otomatis bebas dipakai. Animasi Fox membutuhkan atribusi walau model Fox sendiri CC0.')
doc.save(OUTPUT)
print(OUTPUT)
